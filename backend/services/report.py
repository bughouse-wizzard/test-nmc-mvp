"""
Report generation service for creating XLSX reports from search results.
"""
import io
from typing import List, Dict, Any, Optional
from uuid import UUID
from datetime import datetime
from sqlalchemy.orm import Session

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

from ..models import SearchRequest, ContractResult, SpecComparisonRow, SearchStatus, MatchType


class ReportGenerator:
    """Service for generating XLSX reports from search results."""
    
    def __init__(self, db: Session):
        self.db = db
        
    def generate_search_report(self, search_id: str) -> io.BytesIO:
        """
        Generate an XLSX report for a search.
        
        Args:
            search_id: ID of the search request
            
        Returns:
            BytesIO object containing the XLSX file
        """
        # Get search request
        search_request = self.db.query(SearchRequest).filter(SearchRequest.id == search_id).first()
        if not search_request:
            raise ValueError(f"Search request with ID {search_id} not found")
        
        # Get all contract results for this search
        contract_results = self.db.query(ContractResult)\
            .filter(ContractResult.search_id == search_id)\
            .order_by(ContractResult.created_at.desc())\
            .all()
        
        # Get spec comparison rows for contracts that have them
        contract_ids = [cr.id for cr in contract_results]
        spec_comparisons = {}
        if contract_ids:
            spec_rows = self.db.query(SpecComparisonRow)\
                .filter(SpecComparisonRow.contract_result_id.in_(contract_ids))\
                .all()
            
            # Group by contract result ID
            for row in spec_rows:
                if row.contract_result_id not in spec_comparisons:
                    spec_comparisons[row.contract_result_id] = []
                spec_comparisons[row.contract_result_id].append(row)
        
        # Create workbook
        wb = Workbook()
        
        # Remove default sheet
        if "Sheet" in wb.sheetnames:
            default_sheet = wb["Sheet"]
            wb.remove(default_sheet)
        
        # Create summary sheet
        self._create_summary_sheet(wb, search_request, contract_results)
        
        # Create detailed sheet
        self._create_detailed_sheet(wb, contract_results)
        
        # Create comparison sheet (only for selected/accepted contracts)
        selected_contracts = [cr for cr in contract_results if cr.accepted_for_nmc]
        if selected_contracts and spec_comparisons:
            self._create_comparison_sheet(wb, selected_contracts, spec_comparisons)
        
        # Save to bytes buffer
        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        
        return buffer
    
    def _create_summary_sheet(self, wb: Workbook, search_request: SearchRequest, 
                             contract_results: List[ContractResult]) -> None:
        """
        Create summary sheet with search parameters and NMCK.
        
        Args:
            wb: Workbook object
            search_request: Search request object
            contract_results: List of contract results
        """
        ws = wb.create_sheet(title="Сводка")
        
        # Define styles
        header_font = Font(bold=True, size=12)
        header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
        header_font_color = Font(bold=True, color="FFFFFF", size=11)
        cell_fill = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
        border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        
        # Title
        ws.merge_cells('A1:D1')
        title_cell = ws['A1']
        title_cell.value = f"Отчет по поиску: {search_request.object_name}"
        title_cell.font = Font(bold=True, size=14)
        title_cell.alignment = Alignment(horizontal='center')
        
        # Search parameters section
        ws['A3'] = "Параметры поиска"
        ws['A3'].font = header_font
        
        search_params = [
            ("ID поиска", search_request.id),
            ("Дата создания", search_request.created_at.strftime("%d.%m.%Y %H:%M")),
            ("Статус", search_request.status.value),
            ("Объект закупки", search_request.object_name),
            ("Код КТРУ", search_request.ktru_code),
            ("Код ОКПД2", search_request.okpd2_code or "Не указан"),
            ("Регион заказчика", search_request.customer_region),
            ("Закон", search_request.law),
            ("Период с", search_request.date_from.strftime("%d.%m.%Y")),
            ("Период по", search_request.date_to.strftime("%d.%m.%Y")),
            ("Статусы исполнения", ", ".join(search_request.execution_statuses)),
            ("Лимит контрактов", search_request.limit_contracts),
            ("Источник данных", search_request.input_source.value),
        ]
        
        for i, (param, value) in enumerate(search_params, start=4):
            ws[f'A{i}'] = param
            ws[f'B{i}'] = value
            ws[f'A{i}'].font = Font(bold=True)
            ws[f'A{i}'].fill = cell_fill
            ws[f'B{i}'].fill = cell_fill
            ws[f'A{i}'].border = border
            ws[f'B{i}'].border = border
        
        # Results section
        results_row = len(search_params) + 6
        ws[f'A{results_row}'] = "Результаты расчета НМЦК"
        ws[f'A{results_row}'].font = header_font
        
        # Calculate statistics
        total_contracts = len(contract_results)
        accepted_contracts = len([cr for cr in contract_results if cr.accepted_for_nmc])
        identical_matches = len([cr for cr in contract_results if cr.match_type == MatchType.IDENTICAL])
        homogeneous_matches = len([cr for cr in contract_results if cr.match_type == MatchType.HOMOGENEOUS])
        no_matches = len([cr for cr in contract_results if cr.match_type == MatchType.NO_MATCH])
        
        results_data = [
            ("Всего найдено контрактов", search_request.found_total or 0),
            ("Обработано контрактов", search_request.processed_count or 0),
            ("Включено в отчет", total_contracts),
            ("Принято для расчета НМЦК", accepted_contracts),
            ("Идентичные совпадения", identical_matches),
            ("Однородные товары", homogeneous_matches),
            ("Нет совпадения", no_matches),
            ("Расчетное значение НМЦК", f"{search_request.nmc_value:,.2f} RUB" if search_request.nmc_value else "Не рассчитано"),
            ("Время выполнения", f"{search_request.runtime_ms} мс" if search_request.runtime_ms else "Не измерено"),
        ]
        
        for i, (label, value) in enumerate(results_data, start=results_row + 2):
            ws[f'A{i}'] = label
            ws[f'B{i}'] = value
            ws[f'A{i}'].font = Font(bold=True)
            ws[f'A{i}'].fill = cell_fill
            ws[f'B{i}'].fill = cell_fill
            ws[f'A{i}'].border = border
            ws[f'B{i}'].border = border
            
            # Highlight NMCK value
            if label == "Расчетное значение НМЦК" and search_request.nmc_value:
                ws[f'B{i}'].font = Font(bold=True, color="FF0000", size=12)
        
        # Adjust column widths
        ws.column_dimensions['A'].width = 30
        ws.column_dimensions['B'].width = 40
    
    def _create_detailed_sheet(self, wb: Workbook, contract_results: List[ContractResult]) -> None:
        """
        Create detailed sheet with list of all analyzed contracts.
        
        Args:
            wb: Workbook object
            contract_results: List of contract results
        """
        ws = wb.create_sheet(title="Детализация")
        
        # Define styles
        header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF", size=11)
        cell_fill_even = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
        cell_fill_odd = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
        border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        center_alignment = Alignment(horizontal='center', vertical='center')
        
        # Headers
        headers = [
            "№",
            "Реестровый номер",
            "Дата подписания",
            "Тип совпадения",
            "Цена за единицу",
            "Валюта",
            "Оценка ИИ",
            "Производитель (целевой)",
            "Производитель (найденный)",
            "Совпадение производителя",
            "Контракт 2025+",
            "Принят для НМЦК",
            "Ссылка на контракт"
        ]
        
        for col, header in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = center_alignment
            cell.border = border
        
        # Data rows
        for row_idx, contract in enumerate(contract_results, start=2):
            # Alternate row colors
            fill = cell_fill_even if row_idx % 2 == 0 else cell_fill_odd
            
            # Match type color coding
            match_type_color = {
                MatchType.IDENTICAL: "00B050",  # Green
                MatchType.HOMOGENEOUS: "FFC000",  # Yellow
                MatchType.NO_MATCH: "FF0000",    # Red
            }.get(contract.match_type, "000000")
            
            match_type_fill = PatternFill(start_color=match_type_color, end_color=match_type_color, fill_type="solid")
            
            # Accepted for NMCK color coding
            accepted_fill = PatternFill(start_color="00B050", end_color="00B050", fill_type="solid") if contract.accepted_for_nmc \
                          else PatternFill(start_color="FF0000", end_color="FF0000", fill_type="solid")
            
            # Manufacturer match color coding
            manufacturer_fill = None
            if contract.manufacturer_match is True:
                manufacturer_fill = PatternFill(start_color="00B050", end_color="00B050", fill_type="solid")
            elif contract.manufacturer_match is False:
                manufacturer_fill = PatternFill(start_color="FF0000", end_color="FF0000", fill_type="solid")
            
            # Row data
            row_data = [
                row_idx - 1,  # Sequence number
                contract.reestr_number,
                contract.sign_date.strftime("%d.%m.%Y") if contract.sign_date else "",
                contract.match_type.value,
                f"{contract.unit_price:,.2f}" if contract.unit_price else "",
                contract.currency,
                contract.ai_score,
                contract.manufacturer_target or "",
                contract.manufacturer_found or "",
                "Да" if contract.manufacturer_match else "Нет" if contract.manufacturer_match is False else "Не указано",
                "Да" if contract.is_2025_plus else "Нет",
                "Да" if contract.accepted_for_nmc else "Нет",
                contract.contract_url
            ]
            
            for col_idx, value in enumerate(row_data, start=1):
                cell = ws.cell(row=row_idx, column=col_idx, value=value)
                cell.fill = fill
                cell.border = border
                cell.alignment = center_alignment
                
                # Apply special formatting
                if col_idx == 4:  # Match type column
                    cell.fill = match_type_fill
                    cell.font = Font(bold=True, color="FFFFFF")
                elif col_idx == 12:  # Accepted for NMCK column
                    cell.fill = accepted_fill
                    cell.font = Font(bold=True, color="FFFFFF")
                elif col_idx == 10 and manufacturer_fill:  # Manufacturer match column
                    cell.fill = manufacturer_fill
                    cell.font = Font(bold=True, color="FFFFFF")
        
        # Adjust column widths
        column_widths = [5, 20, 15, 15, 15, 10, 10, 25, 25, 20, 15, 15, 40]
        for col_idx, width in enumerate(column_widths, start=1):
            ws.column_dimensions[get_column_letter(col_idx)].width = width
    
    def _create_comparison_sheet(self, wb: Workbook, selected_contracts: List[ContractResult],
                                spec_comparisons: Dict[str, List[SpecComparisonRow]]) -> None:
        """
        Create comparison sheet with detailed spec comparison.
        
        Args:
            wb: Workbook object
            selected_contracts: List of selected contract results
            spec_comparisons: Dictionary mapping contract_id to spec comparison rows
        """
        ws = wb.create_sheet(title="Сравнение характеристик")
        
        # Define styles
        header_fill = PatternFill(start_color="7030A0", end_color="7030A0", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF", size=11)
        contract_header_fill = PatternFill(start_color="4F81BD", end_color="4F81BD", fill_type="solid")
        contract_header_font = Font(bold=True, color="FFFFFF", size=10)
        cell_fill_even = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
        cell_fill_odd = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
        border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        
        # Match status color coding
        match_status_fills = {
            "MATCH": PatternFill(start_color="00B050", end_color="00B050", fill_type="solid"),
            "DIFF": PatternFill(start_color="FF0000", end_color="FF0000", fill_type="solid"),
            "UNKNOWN": PatternFill(start_color="FFC000", end_color="FFC000", fill_type="solid"),
        }
        
        # Title
        ws.merge_cells('A1:E1')
        title_cell = ws['A1']
        title_cell.value = "Детальное сравнение характеристик для контрактов, принятых в расчет НМЦК"
        title_cell.font = Font(bold=True, size=12)
        title_cell.alignment = Alignment(horizontal='center')
        
        # Create comparison table for each contract
        current_row = 3
        
        for contract in selected_contracts:
            if contract.id not in spec_comparisons:
                continue
                
            spec_rows = spec_comparisons[contract.id]
            if not spec_rows:
                continue
            
            # Contract header
            ws.merge_cells(f'A{current_row}:E{current_row}')
            contract_cell = ws.cell(row=current_row, column=1, 
                                   value=f"Контракт: {contract.reestr_number} - {contract.match_type.value}")
            contract_cell.fill = contract_header_fill
            contract_cell.font = contract_header_font
            contract_cell.alignment = Alignment(horizontal='center')
            current_row += 1
            
            # Comparison headers
            headers = ["Характеристика", "Целевое значение", "Фактическое значение", "Статус совпадения", "Вес"]
            for col, header in enumerate(headers, start=1):
                cell = ws.cell(row=current_row, column=col, value=header)
                cell.fill = header_fill
                cell.font = header_font
                cell.border = border
                cell.alignment = Alignment(horizontal='center', vertical='center')
            
            current_row += 1
            
            # Spec comparison rows
            for spec_idx, spec_row in enumerate(spec_rows):
                fill = cell_fill_even if spec_idx % 2 == 0 else cell_fill_odd
                
                row_data = [
                    spec_row.name,
                    spec_row.target_value or "",
                    spec_row.actual_value or "",
                    spec_row.match_status.value,
                    spec_row.weight
                ]
                
                for col, value in enumerate(row_data, start=1):
                    cell = ws.cell(row=current_row, column=col, value=value)
                    cell.fill = fill
                    cell.border = border
                    
                    # Apply match status coloring
                    if col == 4:  # Match status column
                        status_fill = match_status_fills.get(spec_row.match_status.value)
                        if status_fill:
                            cell.fill = status_fill
                            cell.font = Font(bold=True, color="FFFFFF")
                
                current_row += 1
            
            # Add empty row between contracts for better readability
            current_row += 1
        
        # Adjust column widths
        column_widths = [30, 25, 25, 20, 10]
        for col_idx, width in enumerate(column_widths, start=1):
            ws.column_dimensions[get_column_letter(col_idx)].width = width
