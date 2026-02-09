"""
Document Text Extractor Service
Handles extraction of text from various file formats: PDF, DOCX, and HTML.
"""

import os
import logging
from pathlib import Path
from typing import Optional, Tuple
import mimetypes

# PDF processing
try:
    import PyPDF2
    PDF_AVAILABLE = True
except ImportError:
    PDF_AVAILABLE = False
    PyPDF2 = None

try:
    import pdfplumber
    PDFPLUMBER_AVAILABLE = True
except ImportError:
    PDFPLUMBER_AVAILABLE = False
    pdfplumber = None

# DOCX processing
try:
    import docx
    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False
    docx = None

# HTML processing
try:
    from bs4 import BeautifulSoup
    BS4_AVAILABLE = True
except ImportError:
    BS4_AVAILABLE = False
    BeautifulSoup = None

logger = logging.getLogger(__name__)


def extract_text_from_file(filepath: str) -> Tuple[str, bool]:
    """
    Extract text from a file (PDF, DOCX, or HTML).
    
    Args:
        filepath: Path to the file to extract text from
        
    Returns:
        Tuple of (extracted_text, is_scanned_pdf_flag)
        - extracted_text: Cleaned text content
        - is_scanned_pdf_flag: True if PDF appears to be scanned/image-based (text extraction failed)
        
    Raises:
        FileNotFoundError: If file doesn't exist
        ValueError: If file format is not supported or extraction fails
    """
    if not os.path.exists(filepath):
        raise FileNotFoundError(f"File not found: {filepath}")
    
    # Determine file type
    mime_type, _ = mimetypes.guess_type(filepath)
    file_ext = Path(filepath).suffix.lower()
    
    # Try to extract based on file extension and mime type
    if file_ext == '.pdf' or (mime_type and 'pdf' in mime_type):
        return _extract_from_pdf(filepath)
    elif file_ext == '.docx' or (mime_type and 'word' in mime_type or 'docx' in mime_type):
        return _extract_from_docx(filepath)
    elif file_ext in ['.html', '.htm'] or (mime_type and 'html' in mime_type):
        return _extract_from_html(filepath)
    else:
        # Try to determine by content
        with open(filepath, 'rb') as f:
            header = f.read(1024)
            if b'%PDF' in header:
                return _extract_from_pdf(filepath)
            elif b'PK' in header[:4]:  # DOCX is a ZIP file
                return _extract_from_docx(filepath)
            elif b'<!DOCTYPE' in header or b'<html' in header.lower():
                return _extract_from_html(filepath)
            else:
                raise ValueError(f"Unsupported file format: {filepath}")


def _extract_from_pdf(filepath: str) -> Tuple[str, bool]:
    """
    Extract text from PDF file.
    
    Returns:
        Tuple of (extracted_text, is_scanned_pdf_flag)
    """
    if not PDF_AVAILABLE and not PDFPLUMBER_AVAILABLE:
        raise ImportError("PDF processing libraries not available. Install PyPDF2 or pdfplumber.")
    
    extracted_text = ""
    is_scanned = False
    
    # First try with pdfplumber (better for complex PDFs)
    if PDFPLUMBER_AVAILABLE:
        try:
            with pdfplumber.open(filepath) as pdf:
                for page in pdf.pages:
                    page_text = page.extract_text()
                    if page_text:
                        extracted_text += page_text + "\n"
            if extracted_text.strip():
                return _clean_text(extracted_text), False
        except Exception as e:
            logger.debug(f"pdfplumber extraction failed: {e}")
    
    # Fall back to PyPDF2
    if PDF_AVAILABLE:
        try:
            with open(filepath, 'rb') as file:
                pdf_reader = PyPDF2.PdfReader(file)
                for page in pdf_reader.pages:
                    page_text = page.extract_text()
                    if page_text:
                        extracted_text += page_text + "\n"
            
            if extracted_text.strip():
                return _clean_text(extracted_text), False
            else:
                # No text extracted - likely a scanned PDF
                is_scanned = True
                return "", True
        except Exception as e:
            logger.error(f"PyPDF2 extraction failed: {e}")
            raise ValueError(f"Failed to extract text from PDF: {e}")
    
    # If we get here, both methods failed
    is_scanned = True
    return "", True


def _extract_from_docx(filepath: str) -> Tuple[str, bool]:
    """
    Extract text from DOCX file.
    
    Returns:
        Tuple of (extracted_text, is_scanned_pdf_flag) - always False for is_scanned
    """
    if not DOCX_AVAILABLE:
        raise ImportError("python-docx library not available.")
    
    try:
        doc = docx.Document(filepath)
        extracted_text = ""
        
        # Extract text from paragraphs
        for paragraph in doc.paragraphs:
            if paragraph.text.strip():
                extracted_text += paragraph.text + "\n"
        
        # Extract text from tables
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        extracted_text += cell.text + "\n"
        
        return _clean_text(extracted_text), False
    except Exception as e:
        logger.error(f"DOCX extraction failed: {e}")
        raise ValueError(f"Failed to extract text from DOCX: {e}")


def _extract_from_html(filepath: str) -> Tuple[str, bool]:
    """
    Extract text from HTML file.
    
    Returns:
        Tuple of (extracted_text, is_scanned_pdf_flag) - always False for is_scanned
    """
    if not BS4_AVAILABLE:
        raise ImportError("BeautifulSoup4 library not available.")
    
    try:
        with open(filepath, 'r', encoding='utf-8') as f:
            html_content = f.read()
        
        soup = BeautifulSoup(html_content, 'html.parser')
        
        # Remove script and style elements
        for script in soup(["script", "style"]):
            script.decompose()
        
        # Get text
        text = soup.get_text()
        
        # Clean up whitespace
        lines = (line.strip() for line in text.splitlines())
        chunks = (phrase.strip() for line in lines for phrase in line.split("  "))
        text = '\n'.join(chunk for chunk in chunks if chunk)
        
        return _clean_text(text), False
    except UnicodeDecodeError:
        # Try with different encoding
        try:
            with open(filepath, 'r', encoding='latin-1') as f:
                html_content = f.read()
            
            soup = BeautifulSoup(html_content, 'html.parser')
            for script in soup(["script", "style"]):
                script.decompose()
            
            text = soup.get_text()
            lines = (line.strip() for line in text.splitlines())
            chunks = (phrase.strip() for line in lines for phrase in line.split("  "))
            text = '\n'.join(chunk for chunk in chunks if chunk)
            
            return _clean_text(text), False
        except Exception as e:
            logger.error(f"HTML extraction failed with latin-1 encoding: {e}")
            raise ValueError(f"Failed to extract text from HTML: {e}")
    except Exception as e:
        logger.error(f"HTML extraction failed: {e}")
        raise ValueError(f"Failed to extract text from HTML: {e}")


def _clean_text(text: str) -> str:
    """
    Clean extracted text by removing excessive whitespace and normalizing.
    """
    if not text:
        return ""
    
    # Remove leading/trailing whitespace
    text = text.strip()
    
    # Process line by line
    import re
    lines = text.split('\n')
    cleaned_lines = []
    for line in lines:
        # Collapse multiple spaces to single space
        line = re.sub(r'[ \t]+', ' ', line)
        # Strip the line
        line = line.strip()
        if line:  # Keep non-empty lines
            cleaned_lines.append(line)
    
    # Join with newlines and collapse multiple newlines
    text = '\n'.join(cleaned_lines)
    text = re.sub(r'\n\s*\n\s*\n+', '\n\n', text)
    
    return text